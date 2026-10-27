from io import BytesIO
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from .models import AuditLog


class AuditExcelTests(TestCase):
    def test_filtered_all_pages_export_and_permissions(self):
        admin = get_user_model().objects.create_user(username='excel-admin', is_superuser=True)
        member = get_user_model().objects.create_user(username='excel-user', is_staff=True)
        for index in range(12):
            AuditLog.objects.create(company_name='Export Company', actor_name='=SUM(1,2)', action='view',
                resource='Export test', method='GET', status_code=200, result='success', path=f'/test/{index}/')
        AuditLog.objects.create(company_name='Other Company', action='view', resource='Excluded', method='GET', status_code=200, result='success', path='/excluded/')
        url = reverse('audit_logs')
        self.client.force_login(admin)
        response = self.client.get(url, {'company': 'Export Company', 'per_page': 5, 'page': 2, 'export': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml.sheet', response['Content-Type'])
        self.assertIn('attachment;', response['Content-Disposition'])
        book = load_workbook(BytesIO(response.content))
        sheet = book['로그조회']
        self.assertEqual(sheet.max_row, 13)
        self.assertEqual(sheet['E2'].value, '=SUM(1,2)')
        self.assertEqual(sheet['E2'].data_type, 's')
        self.assertEqual(sheet['F2'].value, '조회')
        self.assertEqual(self.client.get(url, {'start': 'invalid', 'export': 'xlsx'}).status_code, 400)
        self.client.force_login(member)
        self.assertEqual(self.client.get(url, {'export': 'xlsx'}).status_code, 403)
