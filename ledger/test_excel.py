from io import BytesIO
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from .models import Voucher
from .excel import MIME
from .test_helpers import AuthenticatedLedgerTestCase
from .tests import journal_data


class ExcelTests(AuthenticatedLedgerTestCase):
    def setUp(self):
        super().setUp()
        self.user.is_staff = True
        self.user.save(update_fields=['is_staff'])
        self.voucher = Voucher.objects.create(owner=self.user, date='2026-10-24', kind='income', entry_type='sales',
            account_master_id='5001', settlement_account_id='1101', counterparty=self.partner, amount=1000, memo='=1+2')
        other = get_user_model().objects.create_user(username='excel-other')
        Voucher.objects.create(owner=other, date='2026-10-24', kind='income', amount=9999, memo='Foreign secret')

    def download(self, name, **params):
        response = self.client.get(reverse(name), dict(export='xlsx', **params))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], MIME)
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        book = load_workbook(BytesIO(response.content))
        self.assertNotIn('Foreign secret', str([list(ws.values) for ws in book]))
        return book

    def test_five_report_exports_and_numeric_cells(self):
        period = dict(start='2026-10-01', end='2026-10-31')
        book = self.download('income_statement', **period)
        self.assertIn(('당기순이익', None, None, 1000), list(book['손익계산서'].values))
        self.download('balance_sheet', as_of='2026-10-31')
        book = self.download('account_ledger', account='1101', **period)
        self.assertEqual(book['계정별원장'].cell(2, 7).value, 1000)
        book = self.download('dashboard', **period)
        self.assertEqual(book['전표 내역'].cell(2, 12).value, '=1+2')
        self.assertEqual(book['전표 내역'].cell(2, 12).data_type, 's')
        self.assertEqual(book['전표 내역'].cell(2, 8).number_format, '#,##0;[Red]-#,##0')
        self.download('partner_ledger', partner=f'master:{self.partner.pk}', **period)

    def test_defaults_invalid_conditions_and_owner_scope(self):
        self.download('income_statement')
        self.download('balance_sheet')
        self.download('dashboard')
        for name, params in [('account_ledger', {}), ('partner_ledger', {}), ('income_statement', {'start': 'bad'}), ('dashboard', {'start': '2026-11-01', 'end': '2026-10-01'})]:
            self.assertEqual(self.client.get(reverse(name), dict(export='xlsx', **params)).status_code, 400)
        book = self.download('dashboard', start='2026-11-01')
        self.assertEqual(book['전표 내역'].max_row, 1)

    def test_journal_lines_export_and_buttons(self):
        self.client.post(reverse('voucher_new'), journal_data())
        book = self.download('dashboard')
        self.assertEqual(book['기타전표 분개'].max_row, 3)
        for name in ['income_statement', 'balance_sheet', 'dashboard', 'account_ledger', 'partner_ledger']:
            response = self.client.get(reverse(name))
            self.assertContains(response, 'Excel 다운로드')
            self.assertContains(response, 'window.print()')

    def test_regular_user_blocked_on_all_report_buttons_and_direct_exports(self):
        self.user.is_staff = False
        self.user.save(update_fields=['is_staff'])
        message = '인쇄 기능과 다운로드기능은 서비스 구독사용자만 사용가능합니다. 회원정보 변경하단의 사용등록을 하세요.'
        for name in ['income_statement', 'balance_sheet', 'dashboard', 'account_ledger', 'partner_ledger']:
            response = self.client.get(reverse(name))
            self.assertContains(response, message, count=2)
            self.assertNotContains(response, 'window.print()')
            self.assertNotContains(response, 'excel-download')
            response = self.client.get(reverse(name), {'export': 'xlsx'})
            self.assertContains(response, message, status_code=403)

    def test_system_admin_can_export_without_staff_flag(self):
        self.user.is_staff = False
        self.user.is_superuser = True
        self.user.save(update_fields=['is_staff', 'is_superuser'])
        self.download('dashboard')
