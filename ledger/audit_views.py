from .pagination import page_size
from django import forms
from django.core.paginator import Paginator
from django.db.models import Q
from django.shortcuts import render
from django.http import HttpResponse
from django.utils import timezone
from django.utils.http import content_disposition_header
from io import BytesIO
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET
from .models import AuditLog
from .audit import ACTIONS, RESULTS


class AuditFilter(forms.Form):
    company = forms.CharField(label='업체명', required=False, max_length=150)
    username = forms.CharField(label='사용자 아이디', required=False, max_length=150)
    start = forms.DateField(label='시작일', required=False, widget=forms.DateInput(attrs={'type': 'date'}))
    end = forms.DateField(label='종료일', required=False, widget=forms.DateInput(attrs={'type': 'date'}))
    action = forms.ChoiceField(label='작업', required=False, choices=[('', '전체')] + ACTIONS)
    result = forms.ChoiceField(label='결과', required=False, choices=[('', '전체')] + RESULTS)

    def clean(self):
        data = super().clean()
        if data.get('start') and data.get('end') and data['start'] > data['end']:
            raise forms.ValidationError('종료일은 시작일 이후여야 합니다.')
        return data


def audit_excel(logs):
    book = Workbook(write_only=True)
    sheet = book.create_sheet('로그조회')
    sheet.freeze_panes = 'A2'
    from openpyxl.utils import get_column_letter
    for index, width in enumerate([14, 12, 26, 24, 24, 14, 28, 22, 16, 14, 45], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    titles = ['일자', '시간', '업체명', '업체 회원아이디', '작업 사용자', '작업', '기능', '대상 번호', '결과', '응답 코드', '경로']
    header = []
    for title in titles:
        cell = WriteOnlyCell(sheet, value=title)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='146B55')
        header.append(cell)
    sheet.append(header)
    for log in logs.iterator(chunk_size=1000):
        when = timezone.localtime(log.occurred_at)
        values = [when.strftime('%Y-%m-%d'), when.strftime('%H:%M:%S'), log.company_name or '미등록',
            log.company_user.username if log.company_user else '—', log.actor_name or '비로그인',
            dict(ACTIONS).get(log.action, log.action), log.resource, log.target or '—',
            dict(RESULTS).get(log.result, log.result), log.status_code, log.path]
        row = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value)
            if isinstance(value, str):
                cell.data_type = 's'
            row.append(cell)
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    response = HttpResponse(stream.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = content_disposition_header(True, f'로그조회_{timezone.localdate():%Y%m%d}.xlsx')
    response['X-Content-Type-Options'] = 'nosniff'
    return response


@never_cache
@require_GET
def audit_logs(request):
    if not request.user.is_superuser:
        return render(request, 'ledger/permission_denied.html', status=403)
    form = AuditFilter(request.GET)
    logs = AuditLog.objects.select_related('company_user')
    if form.is_valid():
        data = form.cleaned_data
        if data.get('company'):
            logs = logs.filter(Q(company_name__icontains=data['company']) | Q(company_user__subscription__company_name__icontains=data['company']))
        if data.get('username'):
            logs = logs.filter(actor_name__icontains=data['username'])
        if data.get('start'):
            logs = logs.filter(occurred_at__date__gte=data['start'])
        if data.get('end'):
            logs = logs.filter(occurred_at__date__lte=data['end'])
        for field in ('action', 'result'):
            if data.get(field):
                logs = logs.filter(**{field: data[field]})
    else:
        logs = logs.none()
    if request.GET.get('export') == 'xlsx':
        if not form.is_valid():
            return HttpResponse('조회 조건을 확인한 뒤 다시 다운로드하세요.', status=400, content_type='text/plain; charset=utf-8')
        return audit_excel(logs)
    query = request.GET.copy()
    query.pop('page', None)
    page = Paginator(logs, page_size(request)).get_page(request.GET.get('page'))
    return render(request, 'ledger/audit_logs.html', {'form': form, 'page': page, 'query': query.urlencode()})
