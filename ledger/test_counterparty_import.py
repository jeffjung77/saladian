from io import BytesIO
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from .test_helpers import AuthenticatedLedgerTestCase
from .models import Counterparty
from .counterparty_import import HEADERS


class CounterpartyImportTests(AuthenticatedLedgerTestCase):
    def setUp(self):
        super().setUp()
        self.user.is_staff = True
        self.user.save()

    def upload_rows(self, rows):
        book = Workbook()
        book.active.title = '거래처'
        book.active.append(HEADERS)
        for row in rows:
            book.active.append(row)
        stream = BytesIO()
        book.save(stream)
        return self.client.post(reverse('counterparty_import'), {'file': SimpleUploadedFile('partners.xlsx', stream.getvalue())})

    def row(self, name='새거래처'):
        return ['2', name, '1234567890', '대표', '', '01012345678', '', '서울', '', '001234', '']

    def test_template_and_success(self):
        response = self.client.get(reverse('counterparty_import_template'))
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(list(next(book['거래처'].values)), HEADERS)
        self.assertEqual(book['거래처']['J2'].number_format, '@')
        response = self.upload_rows([self.row()])
        self.assertRedirects(response, reverse('counterparties'))
        partner = Counterparty.objects.get(owner=self.user, name='새거래처')
        self.assertEqual(partner.business_number, '123-45-67890')
        self.assertEqual(partner.bank_account, '001234')

    def test_error_rolls_back_entire_file(self):
        invalid = self.row('오류')
        invalid[3] = ''
        response = self.upload_rows([self.row(), invalid])
        self.assertContains(response, '3행')
        self.assertFalse(Counterparty.objects.filter(name='새거래처').exists())

    def test_duplicate_and_formula_rejected(self):
        self.assertContains(self.upload_rows([self.row(), self.row()]), '중복')
        formula = self.row()
        formula[1] = '=1+1'
        self.assertContains(self.upload_rows([formula]), '수식')
        self.assertContains(self.upload_rows([self.row(self.partner.name)]), '이미 등록된')

    def test_permission_denied(self):
        self.user.is_staff = False
        self.user.save()
        for route in ['counterparty_import', 'counterparty_import_template']:
            self.assertEqual(self.client.get(reverse(route)).status_code, 403)
