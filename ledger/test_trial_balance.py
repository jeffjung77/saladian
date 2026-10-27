from datetime import date
from io import BytesIO
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from .models import Voucher, JournalLine
from .reports import trial_balance


class TrialBalanceTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user(username='trial-owner', is_staff=True)
        self.other = get_user_model().objects.create_user(username='trial-other')
        self.client.force_login(self.owner)

    def test_gross_movements_net_balances_cutoff_and_export(self):
        for kind, primary, amount in [('income', '5001', 1000), ('expense', '6503', 300)]:
            Voucher.objects.create(owner=self.owner, date=date(2026, 7, 1), kind=kind,
                amount=amount, account_master_id=primary, settlement_account_id='1101')
        Voucher.objects.create(owner=self.other, date=date(2026, 7, 1), kind='income',
            amount=90000, account_master_id='5001', settlement_account_id='1101')
        Voucher.objects.create(owner=self.owner, date=date(2026, 8, 1), kind='income',
            amount=80000, account_master_id='5001', settlement_account_id='1101')
        source = Voucher.objects.filter(owner=self.owner)
        report = trial_balance(date(2026, 7, 31), vouchers=source)
        cash = next(r for r in report['rows'] if r['account'].pk == '1101')
        self.assertEqual([cash['debit'], cash['credit'], cash['debit_balance']], [1000, 300, 700])
        self.assertEqual(report['totals'], dict(debit=1300, credit=1300, debit_balance=1000, credit_balance=1000))
        response = self.client.get(reverse('trial_balance'), {'as_of': '2026-07-31'})
        self.assertContains(response, '합계잔액 시산표')
        self.assertEqual(response.context['report']['totals'], report['totals'])
        response = self.client.get(reverse('trial_balance'), {'as_of': '2026-07-31', 'export': 'xlsx'})
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book['합계잔액 시산표'].cell(5, 2).value, 1300)
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(reverse('trial_balance'), {'export': 'xlsx'}).status_code, 403)
        self.assertEqual(self.client.get(reverse('trial_balance'), {'as_of': 'invalid'}).status_code, 200)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('trial_balance')).status_code, 302)

    def test_journal_closing_and_invalid_vouchers(self):
        journal = Voucher.objects.create(owner=self.owner, date=date(2026, 6, 30), kind='journal', amount=500, is_closing=True)
        JournalLine.objects.create(voucher=journal, row=1, side='debit', account_id='1101', amount=500)
        JournalLine.objects.create(voucher=journal, row=1, side='credit', account_id='3101', amount=500)
        Voucher.objects.create(owner=self.owner, date=date(2026, 7, 1), kind='income', amount=99)
        report = trial_balance(date(2026, 7, 31), vouchers=Voucher.objects.filter(owner=self.owner))
        self.assertEqual(report['included_count'], 1)
        self.assertEqual(len(report['excluded']), 1)
        self.assertEqual(report['totals']['debit'], 500)
        self.assertEqual(report['balance_difference'], 0)
