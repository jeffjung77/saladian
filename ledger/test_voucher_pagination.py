from io import BytesIO
from django.urls import reverse
from openpyxl import load_workbook
from .models import Voucher
from .test_helpers import AuthenticatedLedgerTestCase


class VoucherPaginationTests(AuthenticatedLedgerTestCase):
    def test_pages_filters_totals_and_excel_include_all(self):
        for index in range(23):
            Voucher.objects.create(owner=self.user, date='2026-10-25', kind='income', entry_type='sales',
                account_master_id='5001', settlement_account_id='1101', amount=100, memo=f'row-{index}')
        Voucher.objects.create(owner=self.user, date='2026-09-01', kind='income', amount=500)
        url = reverse('dashboard')
        params = {'start': '2026-10-01', 'end': '2026-10-31'}
        response = self.client.get(url, params)
        self.assertEqual(len(response.context['voucher_page']), 10)
        self.assertEqual(response.context['count'], 23)
        self.assertEqual(response.context['income'], 2300)
        first_ids = {v.pk for v in response.context['voucher_page']}
        self.assertContains(response, 'start=2026-10-01&amp;end=2026-10-31&amp;page=2')
        response = self.client.get(url, dict(params, page=2))
        self.assertEqual(len(response.context['voucher_page']), 10)
        self.assertFalse(first_ids & {v.pk for v in response.context['voucher_page']})
        response = self.client.get(url, dict(params, page=3))
        self.assertEqual(len(response.context['voucher_page']), 3)
        self.user.is_staff = True
        self.user.save()
        response = self.client.get(url, dict(params, page=3, export='xlsx'))
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book['전표 내역'].max_row, 24)

    def test_empty_and_invalid_page(self):
        url = reverse('dashboard')
        for page in ['invalid', '999', '-1']:
            response = self.client.get(url, {'page': page})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(len(response.context['voucher_page']), 0)
