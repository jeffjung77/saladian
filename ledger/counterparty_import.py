from io import BytesIO
from zipfile import ZipFile, BadZipFile
from django import forms
from django.contrib import messages
from django.db import transaction, IntegrityError
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from .forms import CounterpartyForm
from .permissions import management_required

FIELDS = ['kind', 'name', 'business_number', 'representative', 'contact_name', 'phone', 'email', 'address', 'bank_name', 'bank_account', 'memo']
HEADERS = ['거래처구분', '거래처명', '사업자등록번호', '대표자', '담당자', '연락처', '이메일', '주소', '거래은행명', '계좌번호', '비고']


class UploadForm(forms.Form):
    file = forms.FileField(label='거래처 엑셀 파일', widget=forms.FileInput(attrs={'accept': '.xlsx'}))


@management_required
@require_http_methods(['GET'])
def template(request):
    book = Workbook()
    sheet = book.active
    sheet.title = '거래처'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:K1001'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='146B55')
    for column in 'ABCDEFGHIJK':
        sheet.column_dimensions[column].width = 24 if column != 'H' else 45
    for row in sheet.iter_rows(min_row=2, max_row=1001, max_col=11):
        for cell in row:
            cell.number_format = '@'
    choice = DataValidation(type='list', formula1='"1,2,3,4"')
    choice.errorTitle = '거래처구분 확인'
    choice.error = '1, 2, 3, 4 중 선택하세요.'
    choice.showErrorMessage = True
    sheet.add_data_validation(choice)
    choice.add('A2:A1001')
    guide = book.create_sheet('작성안내')
    for line in ['거래처 시트의 2행부터 입력하세요. 한 번에 최대 1,000건입니다.',
                 '필수: 거래처구분, 거래처명, 사업자등록번호, 대표자, 주소',
                 '거래처구분: 1=고객, 2=구매처, 3=직원, 4=기타',
                 '전화번호·사업자등록번호·계좌번호는 앞자리 0이 유지되도록 텍스트로 입력하세요.',
                 '거래처번호는 자동 생성됩니다. 기존 거래처명은 중복 등록할 수 없습니다.',
                 '오류가 하나라도 있으면 전체 파일을 등록하지 않습니다. 수정 후 다시 업로드하세요.',
                 '양식의 제목과 열 순서를 유지하세요. 수식은 사용할 수 없습니다.']:
        guide.append([line])
    guide.column_dimensions['A'].width = 110
    output = BytesIO()
    book.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="counterparty_template.xlsx"'
    return response


def parse_upload(upload, user):
    if not upload.name.lower().endswith('.xlsx') or upload.size > 5 * 1024 * 1024:
        raise ValueError('5MB 이하의 .xlsx 파일을 선택하세요.')
    try:
        with ZipFile(upload) as archive:
            if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError('압축 해제된 엑셀 파일 크기가 너무 큽니다.')
        upload.seek(0)
        book = load_workbook(upload, read_only=True, data_only=False)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError('올바른 엑셀 양식 파일이 아닙니다.') from exc
    try:
        if '거래처' not in book.sheetnames:
            raise ValueError('거래처 시트를 찾을 수 없습니다. 제공된 양식을 사용하세요.')
        rows = book['거래처'].iter_rows()
        header = next(rows, ())
        if [c.value for c in header] != HEADERS:
            raise ValueError('열 제목이나 순서가 다릅니다. 제공된 양식을 사용하세요.')
        instances, errors, names = [], [], set()
        for number, cells in enumerate(rows, 2):
            if not any(c.value is not None and str(c.value).strip() for c in cells):
                continue
            if len(instances) + len(errors) >= 1000:
                raise ValueError('한 번에 최대 1,000건까지 업로드할 수 있습니다.')
            if any(c.data_type == 'f' for c in cells):
                errors.append(f'{number}행: 수식 대신 값을 입력하세요.')
                continue
            values = [str(c.value).strip() if c.value is not None else '' for c in cells]
            form = CounterpartyForm(dict(zip(FIELDS, values)), user=user)
            if not form.is_valid():
                detail = '; '.join(f'{form.fields[key].label}: {", ".join(value)}' for key, value in form.errors.items())
                errors.append(f'{number}행: {detail}')
                continue
            instance = form.save(commit=False)
            if instance.name in names:
                errors.append(f'{number}행: 파일 내에 같은 거래처명이 중복되어 있습니다.')
            else:
                names.add(instance.name)
                instances.append(instance)
        if not instances and not errors:
            raise ValueError('등록할 거래처를 2행부터 입력하세요.')
        return instances, errors
    finally:
        book.close()


@management_required
@require_http_methods(['GET', 'POST'])
def upload(request):
    form = UploadForm(request.POST or None, request.FILES or None)
    errors = []
    if request.method == 'POST' and form.is_valid():
        try:
            instances, errors = parse_upload(form.cleaned_data['file'], request.user)
            if not errors:
                with transaction.atomic():
                    for instance in instances:
                        instance.save()
                messages.success(request, f'거래처 {len(instances)}건을 등록했습니다.')
                return redirect('counterparties')
        except IntegrityError:
            errors = ['중복 거래처가 발견되어 전체 등록을 취소했습니다. 거래처 목록을 확인하세요.']
        except (ValueError, TypeError, EOFError, BadZipFile) as exc:
            errors = [str(exc)]
    return render(request, 'ledger/counterparty_import.html', {'form': form, 'errors': errors})
