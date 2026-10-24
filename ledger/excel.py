"""Excel exports use the same validated, owner-scoped context as the HTML report."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.http import content_disposition_header
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
TITLES = {'income_statement': '손익계산서', 'balance_sheet': '대차대조표',
          'dashboard': '전표집계', 'account_ledger': '계정별원장', 'partner_ledger': '거래처별원장'}


def report_render(request, template, context):
    if request.GET.get('export') != 'xlsx':
        return render(request, template, context)
    if not (request.user.is_staff or request.user.is_superuser):
        return HttpResponse('인쇄 기능과 다운로드기능은 서비스 구독사용자만 사용가능합니다. 회원정보 변경하단의 사용등록을 하세요.', status=403, content_type='text/plain; charset=utf-8')
    kind = request.resolver_match.url_name
    form = context.get('period') if kind == 'dashboard' else context['form']
    if not form.is_bound or not form.is_valid():
        return HttpResponse('조회 조건을 확인하고 먼저 보고서를 조회하세요.', status=400, content_type='text/plain; charset=utf-8')
    book = Workbook()
    book.remove(book.active)

    def sheet(title, rows):
        ws = book.create_sheet(title)
        for row in rows:
            ws.append([str(value) if value is not None and not isinstance(value, (str, int, float, Decimal, date, datetime)) else value for value in row])
        for row in ws:
            for cell in row:
                # User text must remain literal even when it starts with '='.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if isinstance(cell.value, (int, float, Decimal)):
                    cell.number_format = '#,##0;[Red]-#,##0'
                elif isinstance(cell.value, (date, datetime)):
                    cell.number_format = 'yyyy-mm-dd'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='146B55')
        for column in ws.columns:
            width = max((len(str(cell.value or '')) for cell in column), default=10)
            ws.column_dimensions[column[0].column_letter].width = min(55, max(16, width * 1.5))
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    title = TITLES[kind]
    company = getattr(getattr(request.user, 'subscription', None), 'company_name', '') or getattr(getattr(request.user, 'business_profile', None), 'company_name', '') or request.user.username
    sheet('조회정보', [['항목', '내용'], ['보고서', title], ['업체명', company], ['회원아이디', request.user.username],
          ['생성일시', timezone.localtime().strftime('%Y-%m-%d %H:%M:%S')]] + [[form.fields[key].label, value] for key, value in form.cleaned_data.items()])
    report = context.get('report')
    if kind == 'income_statement':
        rows = [['구분', '계정코드', '계정과목', '금액(원)']]
        for section in report['sections']:
            rows.append([section['name'], '', '합계', section['total']])
            rows.extend([section['name'], item['account'].pk, item['account'].name, item['amount']] for item in section['rows'])
        rows.extend([label, '', '', report[key]] for key, label in [('gross_profit', '매출총이익'), ('operating_profit', '영업이익'), ('pretax_profit', '법인세비용차감전순이익'), ('net_profit', '당기순이익')])
        rows.append(['손익대체 결산전표 제외 건수', '', '', report['closing_count']])
        sheet(title, rows)
    elif kind == 'balance_sheet':
        rows = [['구분', '분류', '계정코드', '계정과목', '금액(원)']]
        for section in report['sections']:
            rows.append([section['name'], '', '', '계정 잔액 합계', section['total']])
            for group in section['groups']:
                rows.extend([section['name'], group['name'], item['account'].pk, item['account'].name, item['amount']] for item in group['rows'])
        rows.extend([label, '', '', '', report[key]] for key, label in [('unclosed_profit', '미결산 누적손익'), ('assets', '자산 총계'), ('liabilities', '부채 총계'), ('equity', '자본 총계(미결산 손익 포함)'), ('liabilities_equity', '부채와 자본 총계'), ('difference', '대차 차액')])
        sheet(title, rows)
    elif kind == 'account_ledger':
        sheet('잔액 요약', [['구분', '금액(원)'], ['기초잔액(차변 양수)', report['opening']], ['기간 차변', report['debit']], ['기간 대변', report['credit']], ['기말잔액(차변 양수)', report['closing']]])
        sheet(title, [['일자', '전표번호', '행', '거래처', '상대계정', '적요', '차변', '대변', '잔액(차변 양수)']] +
              [[r['voucher'].date, r['voucher'].pk, r['row'], r['voucher'].partner_label, r['opposite'], r['memo'], r['debit'], r['credit'], r['balance']] for r in report['rows']])
    else:
        sheet('수입 지출 요약', [['구분', '금액(원)'], ['수입', context['income']], ['지출', context['expense']], ['기간 차액', context['balance']]] +
              ([['기초 누계', context['opening']], ['기말 누계', context['closing']], ['안내', '누계는 수입-지출이며 미수금·미지급금 잔액이 아닙니다.']] if kind == 'partner_ledger' else []))
        vouchers = list(context['vouchers']) if kind == 'dashboard' else [r['voucher'] for r in context['rows']] + list(context['journal_vouchers'])
        sheet('전표 내역', [['일자', '전표번호', '입력유형', '구분', '계정코드', '계정과목', '거래처', '공급가액 또는 분개금액', '부가세', '합계(기타전표는 차대변 각)', '부가세 코드', '적요']] +
              [[v.date, v.pk, v.get_entry_type_display(), v.get_kind_display(), v.account_master_id or '', v.account_master.name if v.account_master_id else v.account, v.partner_label, v.amount, v.vat_amount, v.gross_amount, v.vat_display_code, v.memo] for v in vouchers])
        sheet('기타전표 분개', [['일자', '전표번호', '행', '차변 대변', '계정코드', '계정과목', '금액', '라인 적요']] +
              [[v.date, v.pk, line.row, '차변' if line.side == 'debit' else '대변', line.account_id, line.account.name, line.amount, line.memo] for v in vouchers if v.is_journal for line in v.journal_lines.all()])
        if kind == 'dashboard':
            sheet('계정과목별 공급가액', [['구분', '계정코드', '계정과목', '합계']] + [[g['kind'], g['account_master_id'], g['account'], g['total']] for g in context['groups']])
            sheet('차대변 계정별 집계', [['계정코드', '계정과목', '차변', '대변']] + [[g['account_id'], g['account__name'], g['debit'] or 0, g['credit'] or 0] for g in context['journal_groups']])
        else:
            sheet('거래처 누계', [['일자', '전표번호', '누계']] + [[r['voucher'].date, r['voucher'].pk, r['balance']] for r in context['rows']])
    if report:
        sheet('미반영 전표', [['전표번호', '일자', '사유']] + [[r['voucher'].pk, r['voucher'].date, r['reason']] for r in report['excluded']])
    stream = BytesIO()
    book.save(stream)
    response = HttpResponse(stream.getvalue(), content_type=MIME)
    response['Content-Disposition'] = content_disposition_header(True, f'{title}_{timezone.localdate():%Y%m%d}.xlsx')
    response['Cache-Control'] = 'private, no-store'
    response['X-Content-Type-Options'] = 'nosniff'
    return response
